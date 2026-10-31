"""Каталог «Товары и услуги»: PNG из raw/ → JPG 1080×1080 в jpg/ под именами из колонки «Фото (файл)».
python3 build_jpg.py <UM_Yandex_tovary_uslugi_v01.xlsx>   печатает, каких картинок ещё нет."""
import sys, os
import openpyxl
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
RAW, OUT = os.path.join(HERE, "raw"), os.path.join(HERE, "jpg")
BASE_URL = "https://ilyasorokin94-hue.github.io/umnaya-medicina/ym-catalog/"


def rows(xlsx):
    ws = openpyxl.load_workbook(xlsx)["Каталог"]
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        yield dict(zip(head, r))


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    missing = []
    for row in rows(sys.argv[1]):
        jpg = row["Фото (файл)"]
        png = os.path.join(RAW, os.path.splitext(jpg)[0] + ".png")
        if not os.path.exists(png):
            missing.append(f'{row["№"]} {jpg}'); continue
        im = Image.open(png).convert("RGB")
        s = min(im.size)
        im = im.crop(((im.width - s) // 2, (im.height - s) // 2, (im.width + s) // 2, (im.height + s) // 2))
        im.resize((1080, 1080), Image.LANCZOS).save(os.path.join(OUT, jpg), quality=88, optimize=True)
    print("готово:", len(os.listdir(OUT)), "нет:", len(missing))
    for m in missing:
        print("  нет", m)
